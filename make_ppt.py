#!/usr/bin/env python3
"""Génère une présentation PPTX du projet Habits Counting."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Charte ---
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_DK = RGBColor(0x31, 0x2E, 0x81)
DARK = RGBColor(0x1F, 0x23, 0x33)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF5, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENTS = [RGBColor(0xE1, 0x1D, 0x48), RGBColor(0x0E, 0xA5, 0xE9),
           RGBColor(0x16, 0xA3, 0x4A), RGBColor(0xF5, 0x9E, 0x0B)]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
    sp.shadow.inherit = False
    return sp


def textbox(s, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, color, bold, align, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = "Calibri"
    return tb


def bg(s, color):
    rect(s, 0, 0, SW, SH, color)


def title_band(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.5), INDIGO)
    rect(s, 0, Inches(1.5), SW, Emu(40000), RGBColor(0x37, 0x32, 0xC0))
    textbox(s, Inches(0.6), Inches(0.18), Inches(11), Inches(0.4),
            [(kicker.upper(), 13, RGBColor(0xC7, 0xC4, 0xF5), True, PP_ALIGN.LEFT, 0)])
    textbox(s, Inches(0.6), Inches(0.5), Inches(12), Inches(0.95),
            [(title, 30, WHITE, True, PP_ALIGN.LEFT, 0)], anchor=MSO_ANCHOR.MIDDLE)


# ---------- 1. Titre ----------
s = slide()
bg(s, INDIGO)
rect(s, 0, 0, SW, SH, INDIGO)
# bande décorative
rect(s, 0, Inches(5.4), SW, Inches(2.1), INDIGO_DK)
textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2), [
    ("✅  Suivi des habitudes", 50, WHITE, True, PP_ALIGN.LEFT, 6),
    ("Noter chaque jour ses actions, mesurer sa régularité chaque mois",
     22, RGBColor(0xD7, 0xD5, 0xF7), False, PP_ALIGN.LEFT, 0),
])
textbox(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1.2), [
    ("Application web  ·  Node.js + Express + SQLite  ·  Docker", 16, WHITE, True, PP_ALIGN.LEFT, 4),
    ("Présentation du projet — Juin 2026", 14, RGBColor(0xC7, 0xC4, 0xF5), False, PP_ALIGN.LEFT, 0),
])

# ---------- 2. Le besoin ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Le contexte", "Le besoin")
textbox(s, Inches(0.6), Inches(1.9), Inches(12), Inches(0.9), [
    ("Garder une trace de ses bonnes habitudes au quotidien est difficile : "
     "on oublie, et on perd la vision d'ensemble sur la durée.", 18, DARK, False, PP_ALIGN.LEFT, 0)])
needs = [
    ("Noter simplement", "Cocher en un clic les actions réalisées chaque jour"),
    ("Suivre dans le temps", "Naviguer entre les jours, ne rien perdre"),
    ("Récapituler", "Voir chaque mois sa régularité, habitude par habitude"),
]
x = Inches(0.6)
for i, (t, d) in enumerate(needs):
    cx = x + i * Inches(4.15)
    card = rect(s, cx, Inches(3.0), Inches(3.9), Inches(2.6), WHITE)
    rect(s, cx, Inches(3.0), Inches(3.9), Inches(0.12), ACCENTS[i])
    textbox(s, cx + Inches(0.3), Inches(3.35), Inches(3.3), Inches(2.0), [
        (t, 20, INDIGO_DK, True, PP_ALIGN.LEFT, 8),
        (d, 15, GREY, False, PP_ALIGN.LEFT, 0)])

# ---------- 3. La solution ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Vue d'ensemble", "La solution : une appli en 3 onglets")
tabs = [
    ("📅  Jour", "La liste de tes habitudes du jour. Un clic = coché / décoché. "
                 "Boutons précédent / suivant et « Aujourd'hui »."),
    ("📊  Récap mensuel", "Total d'actions cochées, régularité moyenne, et une barre "
                          "de progression par habitude (X jours / mois + %)."),
    ("⚙️  Mes habitudes", "Créer, renommer, choisir une couleur et supprimer "
                          "ses habitudes."),
]
for i, (t, d) in enumerate(tabs):
    y = Inches(2.0) + i * Inches(1.65)
    rect(s, Inches(0.6), y, Inches(12.1), Inches(1.45), WHITE)
    rect(s, Inches(0.6), y, Inches(0.18), Inches(1.45), ACCENTS[i])
    textbox(s, Inches(1.0), y + Inches(0.15), Inches(11.4), Inches(1.2), [
        (t, 20, INDIGO_DK, True, PP_ALIGN.LEFT, 4),
        (d, 15, GREY, False, PP_ALIGN.LEFT, 0)])

# ---------- 4. Fonctionnalités ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Ce que ça fait", "Fonctionnalités clés")
feats = [
    "Gestion complète des habitudes (CRUD)",
    "Couleur personnalisée par habitude",
    "Cochage quotidien instantané",
    "Navigation jour par jour",
    "Récapitulatif mensuel chiffré",
    "Taux de régularité par habitude",
    "Données persistées (base SQLite)",
    "Interface responsive, sans dépendance lourde",
]
for i, f in enumerate(feats):
    col = i % 2
    row = i // 2
    cx = Inches(0.6) + col * Inches(6.2)
    cy = Inches(2.1) + row * Inches(1.15)
    rect(s, cx, cy, Inches(5.9), Inches(0.95), WHITE)
    dot = rect(s, cx + Inches(0.3), cy + Inches(0.33), Inches(0.28), Inches(0.28), INDIGO)
    textbox(s, cx + Inches(0.8), cy + Inches(0.08), Inches(5.0), Inches(0.8),
            [(f, 15, DARK, False, PP_ALIGN.LEFT, 0)], anchor=MSO_ANCHOR.MIDDLE)

# ---------- 5. Architecture ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Comment c'est construit", "Architecture technique")
blocks = [
    ("Frontend", "HTML / CSS / JS purs\nServis en statique\nAucun framework", ACCENTS[1]),
    ("Backend", "Node.js + Express\nAPI REST JSON\nValidation des entrées", INDIGO),
    ("Données", "SQLite (better-sqlite3)\nTables habits & logs\nFichier persistant", ACCENTS[2]),
]
for i, (t, d, c) in enumerate(blocks):
    cx = Inches(0.6) + i * Inches(4.15)
    rect(s, cx, Inches(2.3), Inches(3.9), Inches(2.7), WHITE)
    rect(s, cx, Inches(2.3), Inches(3.9), Inches(0.7), c)
    textbox(s, cx, Inches(2.4), Inches(3.9), Inches(0.5),
            [(t, 19, WHITE, True, PP_ALIGN.CENTER, 0)])
    textbox(s, cx + Inches(0.35), Inches(3.2), Inches(3.2), Inches(1.7),
            [(d, 15, DARK, False, PP_ALIGN.LEFT, 4)])
# flèches
for ax in (Inches(4.5), Inches(8.65)):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, Inches(3.45), Inches(0.45), Inches(0.4))
    a.fill.solid(); a.fill.fore_color.rgb = GREY; a.line.fill.background(); a.shadow.inherit = False
textbox(s, Inches(0.6), Inches(5.3), Inches(12), Inches(0.6),
        [("Le navigateur appelle l'API REST  →  Express lit / écrit dans SQLite  →  réponse JSON",
          15, GREY, False, PP_ALIGN.CENTER, 0)])

# ---------- 6. API & données ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Le cœur technique", "API REST & modèle de données")
# API
rect(s, Inches(0.6), Inches(2.0), Inches(7.3), Inches(4.7), WHITE)
textbox(s, Inches(0.9), Inches(2.15), Inches(6.8), Inches(0.5),
        [("Points d'entrée API", 18, INDIGO_DK, True, PP_ALIGN.LEFT, 0)])
api_rows = [
    "GET    /api/habits            liste",
    "POST   /api/habits            créer",
    "PUT    /api/habits/:id        modifier",
    "DELETE /api/habits/:id        supprimer",
    "GET    /api/logs?date=…       cochés du jour",
    "POST   /api/logs/toggle       cocher/décocher",
    "GET    /api/summary?month=…   récap mensuel",
]
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(6.8), Inches(3.8))
tf = tb.text_frame; tf.word_wrap = True
for i, r in enumerate(api_rows):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(8)
    run = p.add_run(); run.text = r
    run.font.name = "Consolas"; run.font.size = Pt(13); run.font.color.rgb = DARK
# Données
rect(s, Inches(8.1), Inches(2.0), Inches(4.6), Inches(4.7), INDIGO_DK)
textbox(s, Inches(8.4), Inches(2.15), Inches(4.0), Inches(0.5),
        [("Modèle de données", 18, WHITE, True, PP_ALIGN.LEFT, 0)])
textbox(s, Inches(8.4), Inches(2.8), Inches(4.0), Inches(3.7), [
    ("habits", 16, RGBColor(0xC7,0xC4,0xF5), True, PP_ALIGN.LEFT, 2),
    ("id · name · color · sort_order · archived · created_at", 13, WHITE, False, PP_ALIGN.LEFT, 14),
    ("logs", 16, RGBColor(0xC7,0xC4,0xF5), True, PP_ALIGN.LEFT, 2),
    ("id · habit_id · date", 13, WHITE, False, PP_ALIGN.LEFT, 10),
    ("1 ligne = 1 habitude cochée ce jour", 13, RGBColor(0xC7,0xC4,0xF5), False, PP_ALIGN.LEFT, 4),
    ("Contrainte unique (habit_id, date)", 13, RGBColor(0xC7,0xC4,0xF5), False, PP_ALIGN.LEFT, 0),
])

# ---------- 7. Déploiement Docker ----------
s = slide()
bg(s, LIGHT)
title_band(s, "Mise en production", "Déploiement avec Docker")
rect(s, Inches(0.6), Inches(2.0), Inches(6.0), Inches(4.4), WHITE)
textbox(s, Inches(0.9), Inches(2.2), Inches(5.5), Inches(4.0), [
    ("Conteneurisation", 19, INDIGO_DK, True, PP_ALIGN.LEFT, 10),
    ("Dockerfile multi-stage : build des modules natifs puis image runtime légère (slim).", 15, DARK, False, PP_ALIGN.LEFT, 8),
    ("Exécution en utilisateur non-root.", 15, DARK, False, PP_ALIGN.LEFT, 8),
    ("docker-compose avec volume : la base SQLite survit aux redémarrages.", 15, DARK, False, PP_ALIGN.LEFT, 8),
    (".dockerignore pour une image propre.", 15, DARK, False, PP_ALIGN.LEFT, 0),
])
rect(s, Inches(6.9), Inches(2.0), Inches(5.8), Inches(4.4), DARK)
textbox(s, Inches(7.2), Inches(2.2), Inches(5.3), Inches(0.5),
        [("Démarrage en une commande", 16, RGBColor(0x9C,0xF0,0xC0), True, PP_ALIGN.LEFT, 0)])
tb = s.shapes.add_textbox(Inches(7.2), Inches(2.85), Inches(5.3), Inches(3.3))
tf = tb.text_frame; tf.word_wrap = True
code = ["$ docker compose up --build", "", "→ http://localhost:3000", "", "# ou en local", "$ npm install", "$ npm start"]
for i, c in enumerate(code):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    run = p.add_run(); run.text = c
    run.font.name = "Consolas"; run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x9C,0xF0,0xC0) if c.startswith("$") else WHITE

# ---------- 8. Roadmap ----------
s = slide()
bg(s, LIGHT)
title_band(s, "La suite", "Évolutions possibles")
road = [
    ("Séries (streaks)", "Récompenser les jours consécutifs"),
    ("Objectifs", "Cible hebdo / mensuelle par habitude"),
    ("Export CSV", "Récupérer ses données"),
    ("Statistiques annuelles", "Vue d'ensemble sur 12 mois"),
    ("Multi-utilisateurs", "Comptes & authentification"),
    ("CI/CD", "Build automatique de l'image Docker"),
]
for i, (t, d) in enumerate(road):
    col = i % 2; row = i // 2
    cx = Inches(0.6) + col * Inches(6.2)
    cy = Inches(2.1) + row * Inches(1.55)
    rect(s, cx, cy, Inches(5.9), Inches(1.3), WHITE)
    rect(s, cx, cy, Inches(0.15), Inches(1.3), ACCENTS[i % 4])
    textbox(s, cx + Inches(0.4), cy + Inches(0.18), Inches(5.2), Inches(1.0), [
        (t, 17, INDIGO_DK, True, PP_ALIGN.LEFT, 3),
        (d, 14, GREY, False, PP_ALIGN.LEFT, 0)])

# ---------- 9. Merci ----------
s = slide()
bg(s, INDIGO)
rect(s, 0, Inches(5.4), SW, Inches(2.1), INDIGO_DK)
textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.0), [
    ("Merci !", 48, WHITE, True, PP_ALIGN.LEFT, 8),
    ("Un suivi simple, des habitudes durables.", 22, RGBColor(0xD7,0xD5,0xF7), False, PP_ALIGN.LEFT, 0)])
textbox(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1.0), [
    ("Code : github.com/jeromemathieu/habitscounting", 15, WHITE, False, PP_ALIGN.LEFT, 0)])

prs.save("Habits-Counting-Presentation.pptx")
print("OK — slides:", len(prs.slides._sldIdLst))
